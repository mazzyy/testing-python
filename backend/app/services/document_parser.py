"""
Document Parser Service for extracting text from uploaded documents
"""
import io
from typing import Optional, Dict, Any
import PyPDF2
import docx
from app.services.azure_openai import get_azure_openai_service


class DocumentParserService:
    """
    Service for parsing uploaded documents and extracting structured data
    """
    
    def __init__(self):
        """Initialize the document parser"""
        self.ai_service = get_azure_openai_service()
    
    def extract_text_from_pdf(self, file_content: bytes) -> str:
        """
        Extract text from PDF file
        
        Args:
            file_content: Raw bytes of the PDF file
        
        Returns:
            Extracted text string
        """
        try:
            pdf_file = io.BytesIO(file_content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            text = ""
            
            print(f"[FILE] PDF has {len(pdf_reader.pages)} pages")
            
            for i, page in enumerate(pdf_reader.pages):
                page_text = page.extract_text() or ""
                text += page_text + "\n"
                print(f"   Page {i+1}: {len(page_text)} characters")
            
            print(f"[OK] Total extracted: {len(text)} characters")
            return text.strip()
            
        except Exception as e:
            print(f"[ERROR] Error reading PDF: {str(e)}")
            return ""
    
    def extract_text_from_docx(self, file_content: bytes) -> str:
        """
        Extract text from DOCX file
        
        Args:
            file_content: Raw bytes of the DOCX file
        
        Returns:
            Extracted text string
        """
        try:
            docx_file = io.BytesIO(file_content)
            doc = docx.Document(docx_file)
            
            # Extract text from paragraphs
            paragraphs = [paragraph.text for paragraph in doc.paragraphs]
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.append(cell.text)
            
            text = "\n".join(paragraphs)
            print(f"[OK] DOCX extracted: {len(text)} characters")
            return text.strip()
            
        except Exception as e:
            print(f"[ERROR] Error reading DOCX: {str(e)}")
            return ""
    
    def extract_text(self, file_content: bytes, filename: str) -> str:
        """
        Extract text based on file type
        
        Args:
            file_content: Raw bytes of the file
            filename: Name of the file (to determine type)
        
        Returns:
            Extracted text string
        """
        filename_lower = filename.lower()
        
        if filename_lower.endswith('.pdf'):
            return self.extract_text_from_pdf(file_content)
        elif filename_lower.endswith(('.docx', '.doc')):
            return self.extract_text_from_docx(file_content)
        else:
            print(f"[ERROR] Unsupported file format: {filename}")
            return ""
    
    def parse_document(
        self,
        file_content: bytes,
        filename: str,
        doc_type: str = "transcript",
        user_id: Optional[int] = None
    ) -> Dict[str, Any]:
        """
        Parse document and extract structured data using AI
        
        Args:
            file_content: Raw bytes of the file
            filename: Name of the file
            doc_type: Type of document (transcript, cv, degree, language_cert)
            user_id: User ID for token tracking (optional)
        
        Returns:
            Dict with success status, extracted data, and any errors
        """
        result = {
            "success": False,
            "extracted_data": None,
            "raw_text_preview": None,
            "error": None
        }
        
        try:
            # Extract text from document
            print(f"\n[FILE] Processing {doc_type}: {filename}")
            print(f"   File size: {len(file_content)} bytes")
            
            text = self.extract_text(file_content, filename)
            
            if not text or len(text) < 50:
                result["error"] = "Could not extract sufficient text from document. The file may be image-based (scanned) or corrupted."
                return result
            
            # Store preview
            result["raw_text_preview"] = text[:500]
            
            # Parse with AI
            print(f"[AI] Sending {len(text)} characters to Azure OpenAI...")
            try:
                parsed_data = self.ai_service.parse_document(text, doc_type, user_id=user_id)
                print(f"[AI] parse_document returned: {type(parsed_data)} - {parsed_data is not None}", flush=True)
            except Exception as ai_err:
                err_msg = f"{type(ai_err).__name__}: {ai_err}"
                print(f"[ERROR] AI parsing threw exception: {err_msg}", flush=True)
                result["error"] = f"AI parsing failed: {err_msg}"
                return result
            
            if parsed_data:
                result["success"] = True
                result["extracted_data"] = parsed_data
                
                # Log extracted fields
                print("[OK] Successfully parsed document:", flush=True)
                for key, value in parsed_data.items():
                    if value and value != "null" and value != []:
                        print(f"   [OK] {key}: {str(value)[:100]}", flush=True)
            else:
                result["error"] = "Failed to parse document content. Please try again."
                print(f"[WARN] AI returned no data. parsed_data={parsed_data}", flush=True)
            
        except Exception as e:
            result["error"] = f"Error processing document: {str(e)}"
            print(f"[ERROR] {result['error']}", flush=True)
        
        return result
    
    def merge_parsed_data(
        self,
        existing_data: Dict[str, Any],
        new_data: Dict[str, Any]
    ) -> Dict[str, Any]:
        """
        Merge new parsed data into existing data (non-destructive)
        
        Args:
            existing_data: Current profile data
            new_data: Newly parsed data
        
        Returns:
            Merged data dictionary
        """
        merged = existing_data.copy()
        
        for key, value in new_data.items():
            # Only update if new value is not empty/null
            if value and value != "null" and value != []:
                # Don't overwrite existing values unless they're empty
                if key not in merged or not merged[key]:
                    merged[key] = value
                # For lists, merge them
                elif isinstance(value, list) and isinstance(merged.get(key), list):
                    merged[key] = list(set(merged[key] + value))
        
        return merged


# Singleton instance
_document_parser_service: Optional[DocumentParserService] = None


def get_document_parser_service() -> DocumentParserService:
    """Get or create document parser service instance"""
    global _document_parser_service
    if _document_parser_service is None:
        _document_parser_service = DocumentParserService()
    return _document_parser_service
